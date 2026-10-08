"""Import the Georgian beneficiary register without guessing missing source facts."""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.utils import timezone
from django.utils.translation import gettext_lazy as _
from openpyxl import load_workbook

from apps.accounts.roles import is_system_manager
from apps.audit.models import AuditEvent
from apps.audit.services import record_event
from apps.centers.models import Center, SpecialistProfile, StaffProfile
from apps.core.validators import phone_number_validator

from .models import (
    Beneficiary,
    BeneficiaryWorkbookImport,
    CenterServiceOffering,
    EnrollmentCenterPlacement,
    EnrollmentSpecialistAssignment,
    EnrollmentStateEvent,
    ServiceDefinition,
    ServiceEnrollment,
)

HEADERS = (
    ("services", "სერვისის სახე"),
    ("full_name", "ბენეფიციარის სახელი და გვარი"),
    ("personal_id", "ბენეფიციარის პირადი ნომერი"),
    ("address", "ბენეფიციარის მისამართი"),
    ("phone", "საკონტაქტო ნომერი"),
    ("sex", "სქესი"),
    ("source_status", "სტატუსი"),
    ("birth_date", "ბენეფიციარის დაბადების თარიღი"),
    ("source_age", "ასაკი"),
    ("barthel_category", "ბართელის სკალა"),
    ("specialist", "პასუხისმგებელი სპეცილისტის სახელი და გვარი"),
    ("household", "ოჯახური მდგომარეობა"),
    ("start_date", "პროექტში ჩართვის თარიღი"),
    ("application_number", "ჩართვის განაცხადის ნომერი"),
    ("notes", "შენიშვნა"),
    ("end_date", "პროექტიდან გასვლის თარიღი"),
    ("exit_reason", "გასვლის მიზეზი"),
)
SERVICE_CODES = {"შინმოვლა": "HOME-CARE", "საკვების მიწოდება": "FOOD-DELIVERY"}
SOURCE_LABELS = {
    "source_status": _("Status recorded in workbook"),
    "source_age": _("Age recorded in workbook"),
    "barthel_category": _("Barthel category recorded in workbook"),
    "specialist": _("Responsible specialist recorded in workbook"),
    "household": _("Household recorded in workbook"),
    "notes": _("Notes recorded in workbook"),
    "personal_id": _("Personal ID recorded in workbook"),
    "birth_date": _("Birth date recorded in workbook"),
    "phone": _("Phone recorded in workbook"),
    "application_number": _("Application number recorded in workbook"),
    "start_date": _("Enrollment date recorded in workbook"),
    "end_date": _("Exit date recorded in workbook"),
    "sex": _("Sex recorded in workbook"),
}
ISSUE_MESSAGES = {
    "personal_id": _("Personal ID needs verification. The website field was left blank."),
    "birth_date": _("Birth date needs verification. The website field was left blank."),
    "phone": _("Phone needs verification. The website field was left blank."),
    "application_number": _(
        "Application number was stored as a date in Excel. The website field was left blank."
    ),
    "start_date": _("Enrollment date is missing or invalid. Access dates need review."),
    "end_date": _("Exit date is missing or inconsistent. The exit date was left unset."),
    "sex": _("Sex needs verification. The website field was left blank."),
    "specialist": _(
        "No unique active staff account matched this specialist in the selected center."
    ),
}


class WorkbookImportError(ValueError):
    """Only fixed messages and row numbers may appear in this exception."""


def text_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_name(value) -> str:
    return " ".join(text_value(value).split()).casefold()


def source_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        # This register uses US-style slash dates, including days greater than 12.
        for format_string in ("%m/%d/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), format_string).date()
            except ValueError:
                continue
    return None


@dataclass(repr=False)
class WorkbookRow:
    number: int
    source: dict
    person: dict
    services: tuple[str, ...]
    start: date | None
    end: date | None
    status: str
    issues: list[str] = field(default_factory=list)


@dataclass(repr=False)
class WorkbookPlan:
    digest: str
    rows: list[WorkbookRow]
    service_codes: tuple[str, ...] = tuple(SERVICE_CODES.values())
    snapshot_on: date | None = None
    require_existing_center: bool = False

    def summary(self):
        return {
            "beneficiaries": len(self.rows),
            "enrollments": sum(len(row.services) for row in self.rows),
            "services": dict(Counter(code for row in self.rows for code in row.services)),
            "statuses": dict(Counter(row.status for row in self.rows)),
            "rows_needing_review": sum(bool(row.issues) for row in self.rows),
            "issues": [
                {"row": row.number, "fields": row.issues} for row in self.rows if row.issues
            ],
        }


def read_workbook(path, *, today=None) -> WorkbookPlan:
    today = today or timezone.localdate()
    if Path(path).stat().st_size > 10 * 1024 * 1024:
        raise WorkbookImportError("Workbook exceeds the 10 MB limit.")
    content = Path(path).read_bytes()
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise WorkbookImportError("Expanded workbook exceeds the 50 MB limit.")
        workbook = load_workbook(BytesIO(content), read_only=True, keep_links=False)
    except (BadZipFile, KeyError, ValueError) as exc:
        raise WorkbookImportError("Cannot read the workbook structure.") from exc
    try:
        if len(workbook.worksheets) != 1:
            raise WorkbookImportError("Expected exactly one worksheet.")
        sheet = workbook.worksheets[0]
        if (sheet.max_row or 0) > 5001 or (sheet.max_column or 0) > 40:
            raise WorkbookImportError("Workbook exceeds 5,000 rows or 40 columns.")
        rows = iter(sheet.iter_rows())
        header = next(rows, ())
        if [text_value(cell.value) for cell in header[:17]] != [v for _, v in HEADERS]:
            raise WorkbookImportError("Workbook headers do not match the beneficiary register.")
        result = []
        identities = set()
        for number, cells in enumerate(rows, start=2):
            values = [cell.value for cell in cells]
            if not any(value is not None for value in values):
                continue
            if any(cell.data_type in {"f", "e"} for cell in cells):
                raise WorkbookImportError(f"Row {number}: formulas or error cells need review.")
            if any(value is not None for value in values[17:]):
                raise WorkbookImportError(
                    f"Row {number}: visit data needs a separate dated service-delivery import."
                )
            raw = dict(zip((key for key, _ in HEADERS), values[:17], strict=True))
            source = {key: text_value(value) for key, value in raw.items()}
            labels = [value.strip() for value in source["services"].split(",")]
            if not labels or any(label not in SERVICE_CODES for label in labels):
                raise WorkbookImportError(f"Row {number}: unknown service.")
            if not source["full_name"]:
                raise WorkbookImportError(f"Row {number}: missing beneficiary name.")
            personal_id = source["personal_id"]
            if personal_id and personal_id in identities:
                raise WorkbookImportError(f"Row {number}: duplicate personal ID in source.")
            identities.add(personal_id)
            issues = []
            if not re.fullmatch(r"[0-9]{11}", personal_id):
                issues.append("personal_id")
                personal_id = ""
            birth = source_date(raw["birth_date"])
            if source["birth_date"] and (birth is None or birth > today):
                issues.append("birth_date")
                birth = None
            phone = source["phone"]
            if phone.upper() in {"N/A", "NA"}:
                phone = ""
            elif phone:
                try:
                    phone_number_validator(phone)
                except ValidationError:
                    issues.append("phone")
                    phone = ""
            sex = {"მდედრობითი": "female", "მამრობითი": "male"}.get(source["sex"], "")
            if source["sex"] and not sex:
                issues.append("sex")
            application = source["application_number"]
            if isinstance(raw["application_number"], date):
                issues.append("application_number")
                application = ""
            start = source_date(raw["start_date"])
            end = source_date(raw["end_date"])
            exited = bool(source["exit_reason"] or source["end_date"])
            if start is None:
                issues.append("start_date")
            if exited and (end is None or (start and end <= start) or end > today):
                issues.append("end_date")
                end = None
            status = "exited" if exited else "pending" if not start or start > today else "active"
            person = {
                "full_name": source["full_name"],
                "personal_id": personal_id,
                "address": source["address"],
                "phone": phone,
                "sex": sex,
                "birth_date": birth,
                "application_contract_number": application,
                "enrollment_date": start,
                "exit_date": end,
                "exit_reason": source["exit_reason"],
                "service_type": Beneficiary.ServiceType.OTHER,
                "service_status": "applied" if status == "pending" else status,
            }
            result.append(
                WorkbookRow(
                    number,
                    source,
                    person,
                    tuple(dict.fromkeys(SERVICE_CODES[x] for x in labels)),
                    start,
                    end,
                    status,
                    issues,
                )
            )
        if not result:
            raise WorkbookImportError("Workbook contains no beneficiary rows.")
        return WorkbookPlan(hashlib.sha256(content).hexdigest(), result)
    finally:
        workbook.close()


def source_details(record):
    """Call only after authorizing restricted person access in the view."""
    if record is None:
        return None
    if record.source_values.get("source_format") == "early_intervention_pdf_v1":
        from .early_intervention_import import pdf_source_details

        return pdf_source_details(record)
    if record.source_values.get("source_format") == "early_intervention_docx_v1":
        from .early_intervention_docx_import import docx_source_details

        return docx_source_details(record)
    always = {"source_status", "source_age", "barthel_category", "specialist", "household", "notes"}
    return {
        "row": record.source_row,
        "values": [
            {"label": label, "value": record.source_values.get(key, "")}
            for key, label in SOURCE_LABELS.items()
            if key in always | set(record.review_issues) and record.source_values.get(key)
        ],
        "issues": [ISSUE_MESSAGES[key] for key in record.review_issues if key in ISSUE_MESSAGES],
    }


def validate_destination(plan, *, actor, center_code, center_name, link_specialists):
    if not is_system_manager(actor):
        raise PermissionDenied("Only an active System Manager can import beneficiaries.")
    center = Center.objects.filter(code=center_code).first()
    if center is None and plan.require_existing_center:
        raise WorkbookImportError("The selected existing center could not be found.")
    if center and (center.name != center_name or not center.is_active):
        raise WorkbookImportError("Center code belongs to a different name or inactive center.")
    if not center and Center.objects.filter(name=center_name).exists():
        raise WorkbookImportError("Center name already exists with a different code.")
    existing = {
        record.source_row: record
        for record in BeneficiaryWorkbookImport.objects.filter(
            source_digest=plan.digest
        ).select_related("beneficiary")
    }
    if any(
        record.beneficiary.center_id != getattr(center, "pk", None) for record in existing.values()
    ):
        raise WorkbookImportError("This workbook was already imported into another center.")
    if set(existing) - {row.number for row in plan.rows}:
        raise WorkbookImportError("Previously imported source row count differs. Review required.")
    services = {s.code: s for s in ServiceDefinition.objects.filter(code__in=plan.service_codes)}
    if len(services) != len(plan.service_codes) or any(
        not service.is_active or service.valid_to for service in services.values()
    ):
        raise WorkbookImportError(
            "Required service definitions must exist and be active without an end date."
        )
    if plan.require_existing_center and any(
        service.family != ServiceDefinition.Family.EARLY_INTERVENTION
        for service in services.values()
    ):
        raise WorkbookImportError("The selected service is not an early-intervention service.")
    offerings = {}
    if center:
        for code in services:
            candidates = list(
                CenterServiceOffering.objects.filter(center=center, service=services[code])
            )
            if len(candidates) > 1 or any(
                not item.is_active or item.valid_to for item in candidates
            ):
                raise WorkbookImportError("Existing center offerings need manual date review.")
            offerings[code] = candidates[0] if candidates else None
    specialists = {}
    if link_specialists and center:
        for specialist in SpecialistProfile.objects.filter(
            center_assignments__center=center,
            staff_profile__status=StaffProfile.Status.ACTIVE,
            staff_profile__user__is_active=True,
        ).select_related("staff_profile__user"):
            name = normalized_name(specialist.staff_profile.user.get_full_name())
            specialists.setdefault(name, []).append(specialist)
    matches = {}
    for row in plan.rows:
        if row.number in existing:
            if existing[row.number].source_values != row.source:
                raise WorkbookImportError(
                    f"Row {row.number}: previously imported source differs. Review required."
                )
            continue
        if (
            row.person["personal_id"]
            and Beneficiary.objects.filter(personal_id=row.person["personal_id"]).exists()
        ):
            raise WorkbookImportError(
                f"Row {row.number}: personal ID already exists. Manual reconciliation required."
            )
        matches[row.number] = []
        row.issues = [issue for issue in row.issues if issue != "specialist"]
        if link_specialists and row.source["specialist"]:
            # A register cell may name several jointly responsible specialists.
            # Match each person independently, without inferring a lead from order.
            names = dict.fromkeys(
                normalized_name(name) for name in row.source["specialist"].split(",")
            )
            for name in names:
                candidates = specialists.get(name, [])
                if len(candidates) == 1:
                    matches[row.number].append(candidates[0])
                elif "specialist" not in row.issues:
                    row.issues.append("specialist")
        # Validate normalized fields without writing any personal data during preview.
        person = Beneficiary(**row.person)
        try:
            person.full_clean(exclude=["center"])
        except ValidationError as exc:
            raise WorkbookImportError(
                f"Row {row.number}: normalized person fields failed validation."
            ) from exc
    return center, services, offerings, matches, existing


@transaction.atomic
def apply_workbook(
    plan, *, actor, center_code, center_name, link_specialists, allow_incomplete=False
):
    # Lock catalog rows to serialize imports and repeat all destination checks under the lock.
    list(ServiceDefinition.objects.select_for_update().filter(code__in=plan.service_codes))
    center, services, offerings, matches, existing = validate_destination(
        plan,
        actor=actor,
        center_code=center_code,
        center_name=center_name,
        link_specialists=link_specialists,
    )
    if any(row.issues for row in plan.rows) and not allow_incomplete:
        raise WorkbookImportError("Review issues require explicit --allow-incomplete approval.")
    new_rows = [row for row in plan.rows if row.number not in existing]
    if not new_rows:
        return {
            "created": 0,
            "skipped": len(existing),
            "enrollments_created": 0,
            "assignments_created": 0,
        }
    if center is None:
        center = Center.objects.create(code=center_code, name=center_name)
        audit_create(actor, center, center)
    for code, service in services.items():
        starts = [
            row.start or plan.snapshot_on
            for row in new_rows
            if code in row.services and (row.start or plan.snapshot_on)
        ]
        if not any(code in row.services for row in new_rows):
            continue
        first = min(starts) if starts else None
        if first and service.valid_from and first < service.valid_from:
            service.valid_from = first
            service.save(update_fields=["valid_from", "updated_at"])
            record_event(
                actor=actor,
                center=center,
                event_type=AuditEvent.EventType.UPDATE,
                target_type="ServiceDefinition",
                target_id=service.pk,
            )
        offering = offerings.get(code)
        if offering is None:
            offering = CenterServiceOffering.objects.create(
                center=center, service=service, valid_from=first
            )
            audit_create(actor, center, offering)
        elif first and offering.valid_from and first < offering.valid_from:
            offering.valid_from = first
            offering.save(update_fields=["valid_from", "updated_at"])
            record_event(
                actor=actor,
                center=center,
                event_type=AuditEvent.EventType.UPDATE,
                target_type="CenterServiceOffering",
                target_id=offering.pk,
            )
        offerings[code] = offering
    enrollment_count = assignment_count = 0
    for row in new_rows:
        try:
            beneficiary = Beneficiary.objects.create(
                center=center,
                **row.person,
            )
            audit_create(actor, center, beneficiary)
            source_record = BeneficiaryWorkbookImport.objects.create(
                beneficiary=beneficiary,
                source_digest=plan.digest,
                source_row=row.number,
                source_values=row.source,
                review_issues=row.issues,
            )
            audit_create(actor, center, source_record)
            for code in row.services:
                incomplete = row.start is None or (row.status == "exited" and row.end is None)
                suffix = {
                    "HOME-CARE": "HC",
                    "FOOD-DELIVERY": "FD",
                    "EARLY-INTERVENTION": "EI",
                }[code]
                enrollment = ServiceEnrollment.objects.create(
                    beneficiary=beneficiary,
                    service=services[code],
                    episode_code=f"{beneficiary.beneficiary_code}-{suffix}",
                    status=row.status,
                    start_date=row.start,
                    end_date=row.end,
                    application_contract_number=row.person["application_contract_number"],
                    exit_reason=row.source["exit_reason"],
                    legacy_dates_incomplete=incomplete,
                )
                EnrollmentCenterPlacement.objects.create(
                    enrollment=enrollment,
                    center=center,
                    offering=offerings[code],
                    valid_from=row.start or plan.snapshot_on,
                    valid_to=row.end,
                    legacy_dates_incomplete=incomplete,
                )
                # A snapshot does not establish the dates of earlier lifecycle transitions.
                EnrollmentStateEvent.objects.create(
                    enrollment=enrollment,
                    kind=EnrollmentStateEvent.Kind.LEGACY_IMPORT,
                    new_state=row.status,
                    effective_date=(
                        row.end if row.status == "exited" else row.start or plan.snapshot_on
                    ),
                    actor=actor,
                )
                for specialist in matches[row.number]:
                    # Current responsibility starts now, never infer historic staff access.
                    assignment_start = (
                        max(row.start, timezone.localdate())
                        if row.start and row.status != "exited"
                        else None
                    )
                    EnrollmentSpecialistAssignment.objects.create(
                        enrollment=enrollment,
                        specialist=specialist,
                        valid_from=assignment_start,
                        valid_to=row.end,
                        legacy_dates_incomplete=assignment_start is None,
                    )
                    assignment_count += 1
                audit_create(actor, center, enrollment)
                enrollment_count += 1
        except (ValidationError, ValueError) as exc:
            raise WorkbookImportError(
                f"Row {row.number}: import validation failed. All changes rolled back."
            ) from exc
    return {
        "created": len(new_rows),
        "skipped": len(existing),
        "enrollments_created": enrollment_count,
        "assignments_created": assignment_count,
    }


def audit_create(actor, center, obj):
    record_event(
        actor=actor,
        center=center,
        event_type=AuditEvent.EventType.CREATE,
        target_type=type(obj).__name__,
        target_id=obj.pk,
    )
